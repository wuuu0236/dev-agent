"""docx 解析测试。

背景：`parse_docx` 曾经只读 `doc.paragraphs`，**doc.tables 的内容 100% 静默丢失**
（段落照样保留、文档照样标 ready，用户毫无感知）。而 Word 里信息密度最高的
往往就是表格。本文件把这个 bug 钉死，并把表格序列化的各项契约固化下来。

实测基线（修复前，2×2 + 3×3 两张表 / 7 个关键值）：
    返回块 1 个，长度 51，7 个关键值全部 False
"""
from docx import Document

from src.parser import _table_to_text, parse_docx, parse_file


def _save(tmp_path, build, name="t.docx"):
    """构造 docx：build(doc) 里往 doc 塞内容。"""
    doc = Document()
    build(doc)
    path = tmp_path / name
    doc.save(str(path))
    return str(path)


def _text_of(path):
    """parse_docx 输出的全部文本（本解析器恒定返回 1 个块）。"""
    blocks = parse_docx(path)
    return "\n".join(b["text"] for b in blocks)


# ---------------------------------------------------------------- 回归：内容不丢


def test_table_content_not_lost(tmp_path):
    """🔴 回归测试：表格内容不能丢（这是本文件存在的理由）。"""
    def build(doc):
        doc.add_paragraph("这是第一段正文。")
        t = doc.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "基金代码"
        t.cell(0, 1).text = "年化收益"
        t.cell(1, 0).text = "025490"
        t.cell(1, 1).text = "12.5%"
        doc.add_paragraph("这是第二段正文。")

    text = _text_of(_save(tmp_path, build))
    for key in ["基金代码", "年化收益", "025490", "12.5%"]:
        assert key in text, f"表格内容 {key!r} 丢失了"


def test_multiple_tables_all_kept(tmp_path):
    """多张表格都要保留（曾经一张都不留）。"""
    def build(doc):
        doc.add_paragraph("正文开头。")
        t1 = doc.add_table(rows=2, cols=2)
        t1.cell(0, 0).text = "环境"
        t1.cell(0, 1).text = "Python"
        t1.cell(1, 0).text = "生产"
        t1.cell(1, 1).text = "3.11"
        doc.add_paragraph("正文中间。")
        t2 = doc.add_table(rows=2, cols=2)
        t2.cell(0, 0).text = "库"
        t2.cell(0, 1).text = "用途"
        t2.cell(1, 0).text = "Chroma"
        t2.cell(1, 1).text = "向量存储"

    text = _text_of(_save(tmp_path, build))
    for key in ["生产", "3.11", "Chroma", "向量存储"]:
        assert key in text


# ---------------------------------------------------------------- 阅读顺序


def test_reading_order_preserved(tmp_path):
    """表格必须在段落的真实位置上，不能被整体挪到文末。"""
    def build(doc):
        doc.add_paragraph("AA段落在表格之前。")
        t = doc.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "表头甲"
        t.cell(0, 1).text = "表头乙"
        t.cell(1, 0).text = "MM表格数据"
        t.cell(1, 1).text = "NN表格数据"
        doc.add_paragraph("ZZ段落在表格之后。")

    text = _text_of(_save(tmp_path, build))
    assert text.index("AA段落") < text.index("MM表格") < text.index("ZZ段落")


def test_table_between_paragraphs_keeps_blank_line(tmp_path):
    """块与块之间用空行分隔——单换行会被清洗层的断行合并逻辑粘掉。"""
    def build(doc):
        doc.add_paragraph("段落甲")
        t = doc.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "列一"
        t.cell(0, 1).text = "列二"
        t.cell(1, 0).text = "值一"
        t.cell(1, 1).text = "值二"

    text = _text_of(_save(tmp_path, build))
    assert "\n\n" in text


# ---------------------------------------------------------------- 序列化形式


def test_row_is_self_contained(tmp_path):
    """每行都要「列名：值」配对——这是选行式而非 Markdown 表格的全部理由。"""
    def build(doc):
        t = doc.add_table(rows=2, cols=3)
        for i, h in enumerate(["模型", "维度", "场景"]):
            t.cell(0, i).text = h
        t.cell(1, 0).text = "bge-large-zh"
        t.cell(1, 1).text = "1024"
        t.cell(1, 2).text = "中文检索"

    text = _text_of(_save(tmp_path, build))
    assert "模型：bge-large-zh" in text
    assert "维度：1024" in text
    assert "场景：中文检索" in text


def test_row_starts_with_dash_prefix(tmp_path):
    """每行必须以 `- ` 开头：清洗层把它当"不合并"标记，行结构才不会被吞。

    这是与 cleaner._NO_MERGE_PREFIX 的**隐式契约**，所以要有测试守着。
    """
    def build(doc):
        t = doc.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "代码"
        t.cell(0, 1).text = "收益"
        t.cell(1, 0).text = "025490"
        t.cell(1, 1).text = "12.5%"

    text = _text_of(_save(tmp_path, build))
    assert "- 代码：025490" in text


def test_no_markdown_pipe_syntax(tmp_path):
    """不输出 Markdown 表格语法（`|` 不携带语义，还会稀释嵌入向量）。"""
    def build(doc):
        t = doc.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "甲"
        t.cell(0, 1).text = "乙"
        t.cell(1, 0).text = "丙"
        t.cell(1, 1).text = "丁"

    text = _text_of(_save(tmp_path, build))
    assert "|" not in text


# ---------------------------------------------------------------- 合并单元格


def test_horizontal_merge_not_duplicated(tmp_path):
    """横向合并：同一个 <w:tc> 在多个列位重复出现，不能把同一个值挂到两个列名下。"""
    def build(doc):
        t = doc.add_table(rows=3, cols=3)
        for i, h in enumerate(["配置项", "值", "说明"]):
            t.cell(0, i).text = h
        t.cell(1, 0).text = "超时时间"
        t.cell(1, 1).text = "60 秒"
        t.cell(1, 2).text = "超过则重试"
        t.cell(2, 0).merge(t.cell(2, 1))
        t.cell(2, 0).text = "本节无配置"

    text = _text_of(_save(tmp_path, build))
    # 合并值只出现一次，且不会以「值：」的列名重复一次
    assert text.count("本节无配置") == 1
    assert "值：本节无配置" not in text


def test_vertical_merge_repeats_group_value(tmp_path):
    """纵向合并：故意保留重复——组内每行补全所属分组，避免"孤儿行"。"""
    def build(doc):
        t = doc.add_table(rows=3, cols=2)
        t.cell(0, 0).text = "环境"
        t.cell(0, 1).text = "框架"
        t.cell(1, 0).text = "生产"
        t.cell(1, 1).text = "FastAPI"
        t.cell(2, 0).merge(t.cell(1, 0))     # 环境列纵向合并
        t.cell(2, 1).text = "Streamlit"

    text = _text_of(_save(tmp_path, build))
    # 两行都要知道自己在「生产」环境下，不能出现只有「框架：Streamlit」的孤儿行
    assert "环境：生产；框架：FastAPI" in text
    assert "环境：生产；框架：Streamlit" in text


# ---------------------------------------------------------------- 降级与边界


def test_single_row_table_is_data_not_header(tmp_path):
    """只有一行的表没有列名可配，不能把值当列名用掉。"""
    def build(doc):
        t = doc.add_table(rows=1, cols=2)
        t.cell(0, 0).text = "025490"
        t.cell(0, 1).text = "12.5%"

    text = _text_of(_save(tmp_path, build))
    assert "025490" in text and "12.5%" in text


def test_empty_header_falls_back_to_position(tmp_path):
    """表头单元格空缺时退回「第 N 列」，保证值不丢。"""
    def build(doc):
        t = doc.add_table(rows=2, cols=3)
        t.cell(0, 0).text = "名称"
        t.cell(0, 1).text = ""
        t.cell(0, 2).text = "备注"
        t.cell(1, 0).text = "清洗"
        t.cell(1, 1).text = "默认开"
        t.cell(1, 2).text = ""

    text = _text_of(_save(tmp_path, build))
    assert "名称：清洗" in text
    assert "第2列：默认开" in text


def test_empty_cells_skipped(tmp_path):
    """空单元格不产生 `列名：` 这种空配对。"""
    def build(doc):
        t = doc.add_table(rows=2, cols=3)
        t.cell(0, 0).text = "甲"
        t.cell(0, 1).text = "乙"
        t.cell(0, 2).text = "丙"
        t.cell(1, 0).text = "有值"
        t.cell(1, 1).text = ""

    text = _text_of(_save(tmp_path, build))
    assert "乙：" not in text
    assert "丙：" not in text


def test_all_empty_table_produces_no_block(tmp_path):
    """全空表不产生块（否则会塞一堆空块进库）。"""
    def build(doc):
        doc.add_table(rows=2, cols=2)

    assert parse_docx(_save(tmp_path, build)) == []


def test_multi_paragraph_cell_flattened(tmp_path):
    """单元格内多段落要压平，否则会打断「列名：值」配对。"""
    def build(doc):
        t = doc.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "问题"
        t.cell(0, 1).text = "答案"
        t.cell(1, 0).text = "为什么用 RRF\n而不是加权求和"
        t.cell(1, 1).text = "分数量纲不同"

    text = _text_of(_save(tmp_path, build))
    assert "问题：为什么用 RRF 而不是加权求和" in text


def test_table_only_document(tmp_path):
    """整篇只有表格、没有段落，也要正常返回。"""
    def build(doc):
        t = doc.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "键"
        t.cell(0, 1).text = "值"
        t.cell(1, 0).text = "chunk_size"
        t.cell(1, 1).text = "500"

    text = _text_of(_save(tmp_path, build))
    assert "键：chunk_size" in text and "值：500" in text


def test_empty_document_returns_empty(tmp_path):
    """空文档返回空列表。"""
    assert parse_docx(_save(tmp_path, lambda doc: None)) == []


def test_no_table_document_unchanged(tmp_path):
    """没有表格的文档行为不回归：段落照样提取。"""
    def build(doc):
        doc.add_paragraph("第一段。")
        doc.add_paragraph("第二段。")

    text = _text_of(_save(tmp_path, build))
    assert "第一段。" in text and "第二段。" in text


# ---------------------------------------------------------------- 端到端（含清洗层）


def test_survives_cleaning_end_to_end(tmp_path):
    """🔑 端到端：走真实 parse_file（内含清洗）后表格内容与行结构都还在。

    清洗层的 merge_broken_lines 会把「不以句末标点结尾」的行与下一行合并，
    而 `12.5%` 结尾必然被吞 —— 这个测试就是守这道关的。
    """
    def build(doc):
        doc.add_paragraph("这是正文段落。")
        t = doc.add_table(rows=3, cols=2)
        t.cell(0, 0).text = "基金代码"
        t.cell(0, 1).text = "年化收益"
        t.cell(1, 0).text = "025490"
        t.cell(1, 1).text = "12.5%"
        t.cell(2, 0).text = "270042"
        t.cell(2, 1).text = "8.3%"
        doc.add_paragraph("这是收尾段落。")

    text = "\n".join(b["text"] for b in parse_file(_save(tmp_path, build)))

    assert "基金代码：025490；年化收益：12.5%" in text
    assert "基金代码：270042；年化收益：8.3%" in text
    # 两行没有被清洗层粘成一行
    assert "12.5%- " not in text and "12.5% 基金代码" not in text
    # 表格后面的段落也没被粘进表格行
    assert "12.5%这是收尾段落" not in text
    assert "这是收尾段落。" in text


def test_end_to_end_chunking_keeps_table_rows_whole(tmp_path):
    """端到端：分块后表格行仍保持完整（不会被切在「列名：值」中间）。"""
    from src.chunker import chunk_parsed

    def build(doc):
        doc.add_paragraph("以下是模型对比：")
        t = doc.add_table(rows=11, cols=3)
        for i, h in enumerate(["模型", "维度", "场景"]):
            t.cell(0, i).text = h
        for r in range(1, 11):
            t.cell(r, 0).text = f"模型{r}"
            t.cell(r, 1).text = f"{128 * r}"
            t.cell(r, 2).text = f"第{r}类中文语义检索场景"

    chunks = chunk_parsed(parse_file(_save(tmp_path, build)))
    joined = "\n".join(c["content"] for c in chunks)

    assert "模型10" in joined
    # 每一行的「列名：值」三元组必须完整落在同一个 chunk 里
    for c in chunks:
        for line in c["content"].split("\n"):
            if line.startswith("- ") and "模型" in line:
                assert "维度：" in line and "场景：" in line, f"行被切断: {line!r}"


def test_unknown_extension_placeholder(tmp_path):
    """不支持的类型仍返回占位块（向后兼容，不受本次改动影响）。"""
    p = tmp_path / "x.xyz"
    p.write_text("whatever", encoding="utf-8")
    blocks = parse_file(str(p))
    assert len(blocks) == 1
    assert "暂不支持" in blocks[0]["text"]


# ---------------------------------------------------------------- 直测序列化函数


def test_table_to_text_empty_table():
    doc = Document()
    assert _table_to_text(doc.add_table(rows=2, cols=2)) == ""


def test_table_to_text_single_row():
    doc = Document()
    t = doc.add_table(rows=1, cols=3)
    t.cell(0, 0).text = "甲"
    t.cell(0, 1).text = "乙"
    t.cell(0, 2).text = "丙"
    assert _table_to_text(t) == "- 甲；乙；丙"


def test_table_to_text_skips_blank_rows():
    """全空行要跳过（否则会产生空的 `- ` 行）。"""
    doc = Document()
    t = doc.add_table(rows=3, cols=2)
    t.cell(0, 0).text = "键"
    t.cell(0, 1).text = "值"
    # 第 1 行整行留空
    t.cell(2, 0).text = "k"
    t.cell(2, 1).text = "v"
    out = _table_to_text(t)
    assert out == "- 键：k；值：v"
    assert "- \n" not in out + "\n"
